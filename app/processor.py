"""
生成诊断 Excel + 临时文件清理
"""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from .analyzer import run_diagnosis


def _style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws, max_width=36):
    for column in ws.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                max_length = max(max_length, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[letter].width = min(max(max_length + 2, 10), max_width)


def build_excel(std: pd.DataFrame, report: dict, warnings: list[str]) -> bytes:
    wb = Workbook()

    # 1 诊断摘要
    ws = wb.active
    ws.title = "诊断摘要"
    ws.append(["项目", "数值"])
    _style_header(ws)
    for k, v in report.get("overview", {}).items():
        ws.append([k, v])
    ws.append([])
    ws.append(["【行动建议】", ""])
    for i, t in enumerate(report.get("tips", []), 1):
        ws.append([f"建议{i}", t])
    if warnings:
        ws.append([])
        ws.append(["【解析提示】", ""])
        for w in warnings:
            ws.append(["提示", w])
    _auto_width(ws)

    # 2 标准化明细（最多导出 5000 行避免过大）
    ws2 = wb.create_sheet("标准化明细")
    detail = std.head(5000)
    if not detail.empty:
        for r in dataframe_to_rows(detail, index=False, header=True):
            ws2.append(r)
        _style_header(ws2)
        _auto_width(ws2)
    else:
        ws2.append(["无数据"])

    # 3 日趋势
    trend = report.get("trend")
    ws3 = wb.create_sheet("日销售趋势")
    if trend is not None and not trend.empty:
        for r in dataframe_to_rows(trend, index=False, header=True):
            ws3.append(r)
        _style_header(ws3)
        _auto_width(ws3)
    else:
        ws3.append(["无日期字段，无法生成趋势"])

    # 4 商品排行
    pr = report.get("product_rank")
    ws4 = wb.create_sheet("商品排行")
    if pr is not None and not pr.empty:
        for r in dataframe_to_rows(pr.head(100), index=False, header=True):
            ws4.append(r)
        _style_header(ws4)
        _auto_width(ws4)
    else:
        ws4.append(["无商品字段"])

    # 5 渠道 / 地区
    ws5 = wb.create_sheet("渠道与地区")
    ch = report.get("channel_rank")
    ws5.append(["—— 渠道 ——"])
    if ch is not None and not ch.empty:
        for r in dataframe_to_rows(ch, index=False, header=True):
            ws5.append(r)
    else:
        ws5.append(["无渠道字段"])
    ws5.append([])
    ws5.append(["—— 地区 Top30 ——"])
    rg = report.get("region_rank")
    if rg is not None and not rg.empty:
        for r in dataframe_to_rows(rg, index=False, header=True):
            ws5.append(r)
    else:
        ws5.append(["无地区字段"])
    _auto_width(ws5)

    # 6 客户集中度
    ws6 = wb.create_sheet("客户集中度")
    cs = report.get("customer_stats") or {}
    if cs:
        ws6.append(["指标", "数值"])
        _style_header(ws6)
        for k, v in cs.items():
            ws6.append([k, v])
    else:
        ws6.append(["无客户字段，跳过集中度分析"])
    _auto_width(ws6)

    # 7 异常
    ws7 = wb.create_sheet("异常提示")
    an = report.get("anomalies")
    if an is not None and not an.empty:
        for r in dataframe_to_rows(an, index=False, header=True):
            ws7.append(r)
        _style_header(ws7)
        red = PatternFill("solid", fgColor="FFC7CE")
        for row in ws7.iter_rows(min_row=2):
            for cell in row:
                cell.fill = red
        _auto_width(ws7)
    else:
        ws7.append(["未发现明显异常"])
        ws7["A1"].font = Font(color="006600")

    # 8 隐私
    ws8 = wb.create_sheet("隐私说明")
    for line in [
        "本工具隐私承诺",
        "1. 上传文件仅在临时目录处理，完成后立即删除。",
        "2. 不建立账号，不落库，不用于模型训练或第三方共享。",
        "3. 诊断结果仅供经营参考，重要决策请结合业务人工复核。",
        "4. 建议处理完成后立即下载并关闭页面。",
    ]:
        ws8.append([line])
    ws8.column_dimensions["A"].width = 72

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def process_file(path: Path, original_name: str) -> tuple[bytes, dict]:
    std, report, warnings = run_diagnosis(path, original_name)
    excel = build_excel(std, report, warnings)
    summary = {
        "ok": True,
        "overview": report.get("overview", {}),
        "tip_count": len(report.get("tips", [])),
        "anomaly_count": len(report.get("anomalies", [])),
        "warnings": warnings,
        "filename": original_name,
    }
    return excel, summary


def safe_cleanup(paths: list[Path]):
    for p in paths:
        try:
            if p.exists():
                p.unlink()
        except Exception:
            pass
